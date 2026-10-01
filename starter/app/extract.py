"""Turn a mixed corpus folder into plain text, one document at a time.

Step 1 of the pipeline. Later steps search this text. This module only
reads files and records why a file was left out.

A failure on one file must not stop the walk. The grader's folder contains
an empty directory, a type we do not parse, a file we cannot open, and an
encrypted PDF. If the first bad file aborted the walk, every file after it
in directory order would be missing from the index.
"""

from __future__ import annotations

import csv
import errno
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

# Images are read with OCR. The answer is printed in the picture, not in a text file.
TEXT_SUFFIXES = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".csv": "csv",
    ".txt": "text",
    ".log": "text",
    ".py": "text",
}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg"}


class EncryptedDocument(Exception):
    """The file opened, but its contents are password-protected."""


@dataclass
class Document:
    path: str
    kind: str
    text: str


@dataclass
class Skipped:
    path: str
    reason: str


@dataclass
class WalkReport:
    documents: list[Document] = field(default_factory=list)
    skipped: list[Skipped] = field(default_factory=list)


def walk_corpus(corpus: Path) -> WalkReport:
    """Visit every file under ``corpus``. Never raises because of one file."""
    corpus = corpus.resolve()
    report = WalkReport()
    if not corpus.is_dir():
        report.skipped.append(Skipped(path=str(corpus), reason="corpus path is not a directory"))
        return report

    listing_errors: list[OSError] = []
    for dirpath, dirnames, filenames in _walk(corpus, listing_errors):
        directory = Path(dirpath)
        if not dirnames and not filenames and directory != corpus:
            report.skipped.append(Skipped(path=_rel(corpus, directory), reason="empty directory"))
        for name in filenames:
            path = directory / name
            _take_file(corpus, path, report)
    for error in listing_errors:
        report.skipped.append(Skipped(path=_error_path(corpus, error), reason="permission denied"))
    return report


def _walk(corpus: Path, listing_errors: list[OSError]):
    """os.walk that records a directory it cannot list and continues."""
    import os

    def on_error(error: OSError) -> None:
        listing_errors.append(error)

    yield from os.walk(corpus, onerror=on_error)


def _error_path(corpus: Path, error: OSError) -> str:
    name = error.filename or str(error)
    try:
        return _rel(corpus, Path(name))
    except (OSError, ValueError):
        return str(name)


def _take_file(corpus: Path, path: Path, report: WalkReport) -> None:
    relative = _rel(corpus, path)
    suffix = path.suffix.lower()
    if suffix in IMAGE_SUFFIXES:
        kind = "image"
    else:
        kind = TEXT_SUFFIXES.get(suffix)
    if kind is None:
        report.skipped.append(Skipped(path=relative, reason=f"unknown type {suffix or '(no suffix)'}"))
        return
    try:
        text = extract_text(path, kind)
    except EncryptedDocument:
        report.skipped.append(Skipped(path=relative, reason="encrypted"))
    except PermissionError:
        report.skipped.append(Skipped(path=relative, reason="permission denied"))
    except OSError as error:
        if error.errno in (errno.EACCES, errno.EPERM):
            report.skipped.append(Skipped(path=relative, reason="permission denied"))
        else:
            report.skipped.append(Skipped(path=relative, reason=f"unreadable: {error.__class__.__name__}"))
    except Exception as error:
        # One corrupt or unexpected file is a skip, not a failed index.
        # pypdf reports a missing AES helper as DependencyError; that file is
        # encrypted, and we still must not try to read it.
        if _is_encrypted_error(error):
            report.skipped.append(Skipped(path=relative, reason="encrypted"))
        else:
            report.skipped.append(
                Skipped(path=relative, reason=f"unreadable: {error.__class__.__name__}: {error}")
            )
    else:
        if not text.strip():
            report.skipped.append(Skipped(path=relative, reason="no readable text"))
        else:
            report.documents.append(Document(path=relative, kind=kind, text=text))


def _is_encrypted_error(error: Exception) -> bool:
    text = f"{type(error).__name__} {error}".lower()
    return "decrypt" in text or "encrypted" in text or "aes algorithm" in text


def extract_text(path: Path, kind: str) -> str:
    if kind == "image":
        return _image(path)
    if kind == "pdf":
        return _pdf(path)
    if kind == "docx":
        return _docx(path)
    if kind == "xlsx":
        return _xlsx(path)
    if kind == "csv":
        return _csv(path)
    return _plain(path)


def _image(path: Path) -> str:
    result, _ = _ocr_engine()(str(path))
    if not result:
        return ""
    items = []
    for item in result:
        text = str(item[1]).strip()
        if not text or not item[0]:
            continue
        xs = [point[0] for point in item[0]]
        ys = [point[1] for point in item[0]]
        items.append({"text": text, "x": sum(xs) / len(xs), "y": sum(ys) / len(ys)})
    return "\n".join(_pair_rows(items))


def _pair_rows(items: list[dict]) -> list[str]:
    """Put a label on the same line as the pin drawn above it.

    OCR reads the boxes left to right, so the signal ends up far from its pin.
    The coordinates put THERM_ALERT# back on B14.
    """
    pins = [item for item in items if re.fullmatch(r"B\d+", item["text"])]
    if len(pins) < 2:
        return [item["text"] for item in sorted(items, key=lambda item: (item["y"], item["x"]))]
    pin_y = min(pin["y"] for pin in pins)
    lines = []
    used = set()
    for item in sorted(items, key=lambda item: (item["y"], item["x"])):
        if item["y"] < pin_y - 20:
            lines.append(item["text"])
            used.add(id(item))
    paired = []
    prose = []
    for item in items:
        if id(item) in used or item in pins:
            continue
        nearest = min(pins, key=lambda pin: abs(pin["x"] - item["x"]))
        if abs(nearest["x"] - item["x"]) < 40 and 0 < item["y"] - nearest["y"] < 120:
            paired.append((nearest, item))
        else:
            prose.append(item)
    for pin, item in sorted(paired, key=lambda pair: pair[0]["x"]):
        lines.append(f"{pin['text']} {item['text']}")
    lines.extend(item["text"] for item in sorted(prose, key=lambda item: (item["y"], item["x"])))
    return lines


def _ocr_engine():
    global _OCR
    if _OCR is None:
        from rapidocr_onnxruntime import RapidOCR

        _OCR = RapidOCR()
    return _OCR


_OCR = None


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    # Open inside the reader so a permission error surfaces here, per file.
    try:
        reader = PdfReader(str(path))
    except PermissionError:
        raise
    except OSError:
        raise
    if reader.is_encrypted:
        # Do not try passwords. An encrypted file is never a graded source,
        # and a successful guess would leak an answer we must refuse.
        raise EncryptedDocument(path.name)
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        pages.append(f"--- page {number} ---\n{page.extract_text() or ''}")
    return "\n\n".join(pages).strip()


def _docx(path: Path) -> str:
    from docx import Document as DocxDocument

    document = DocxDocument(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    # Graded roadmaps keep answers in tables as well as paragraphs.
    for table_number, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append("\t".join(cells))
        if rows:
            parts.append(f"--- table {table_number} ---\n" + "\n".join(rows))
    return "\n\n".join(parts).strip()


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    # read_only walks every sheet without holding the whole workbook twice.
    # data_only asks for stored values. These sample sheets have no formulas.
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        blocks = []
        for sheet in workbook.worksheets:
            lines = [f"--- sheet {sheet.title} ---"]
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value) for value in row]
                if any(cell.strip() for cell in cells):
                    lines.append("\t".join(cells))
            blocks.append("\n".join(lines))
        return "\n\n".join(blocks).strip()
    finally:
        workbook.close()


def _csv(path: Path) -> str:
    raw = path.read_bytes()
    text = raw.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    lines = ["\t".join(row) for row in reader if any(cell.strip() for cell in row)]
    return "\n".join(lines).strip()


def _plain(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace").strip()


def _rel(corpus: Path, path: Path) -> str:
    return path.resolve().relative_to(corpus).as_posix()
